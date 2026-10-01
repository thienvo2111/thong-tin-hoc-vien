# -*- coding: utf-8 -*-
# Script TAM (co the xoa sau khi dung xong) - remap dia_ban_ma cho final_don_vi_cong_tac_01_so_gddt.xlsx
# dung lai dung logic/bang anh xa da dung trong scripts/t19_remap_dia_danh.py cho file truong.
import re
import unicodedata
from pathlib import Path
from collections import defaultdict

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
BACKUP_DIR = ROOT / "scripts" / "_backup_before_t19"

NEW_SRC = DATA / "Danh-muc-Phuong-xa_moi_34-tinh-thanh-sau-sat-nhap.xlsx"
OLD_TINH_BACKUP = BACKUP_DIR / "final_dia_danh_01_tinh_thanh.xlsx"
OLD_PXA_BACKUP = BACKUP_DIR / "final_dia_danh_02_phuong_xa.xlsx"
SO_GDDT = DATA / "final_don_vi_cong_tac_01_so_gddt.xlsx"

PROVINCE_PREFIXES = ["thanh pho", "tp", "tinh"]
WARD_PREFIXES = ["dac khu", "thi tran", "phuong", "xa"]


def strip_accents(s: str) -> str:
    s = s.replace("đ", "d").replace("Đ", "D")
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return unicodedata.normalize("NFC", s)


def basic_norm(s: str) -> str:
    s = strip_accents(s or "").lower()
    s = s.replace(".", " ").replace("-", " ")
    s = " ".join(s.split())
    return s


def strip_prefix(s: str, prefixes) -> str:
    for p in sorted(prefixes, key=len, reverse=True):
        if s == p:
            return ""
        if s.startswith(p + " "):
            return s[len(p) + 1 :].strip()
    return s


def norm_province(s: str) -> str:
    return strip_prefix(basic_norm(s), PROVINCE_PREFIXES)


def norm_ward(s: str) -> str:
    return strip_prefix(basic_norm(s), WARD_PREFIXES)


def read_rows(path: Path, sheet_index=0, min_row=2):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[sheet_index]]
    rows = list(ws.iter_rows(min_row=min_row, values_only=True))
    wb.close()
    return rows


def write_xlsx(path: Path, header, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    wb.save(path)


def main():
    out_lines = []

    old_tinh_rows = read_rows(OLD_TINH_BACKUP)  # ma, ten, cap, parent_ma (old GSO)
    old_pxa_rows = read_rows(OLD_PXA_BACKUP)  # ma, ten, cap, parent_ma(=old tinh ma)
    so_gddt_rows = read_rows(SO_GDDT)  # ma_don_vi, ten_don_vi, loai_don_vi, dia_ban_ma, don_vi_cha_ma

    out_lines.append(f"old tinh (backup): {len(old_tinh_rows)} dong | old phuong_xa (backup): {len(old_pxa_rows)} dong | so_gddt: {len(so_gddt_rows)} dong")

    # ---- Mapping ten tinh cu(GSO) -> ma BNV moi (giong het logic trong t19_remap_dia_danh.py) ----
    new_wb = openpyxl.load_workbook(NEW_SRC, read_only=True, data_only=True)
    new_ws = new_wb[new_wb.sheetnames[0]]
    new_rows = list(new_ws.iter_rows(min_row=2, values_only=True))
    new_wb.close()

    new_provinces = {}
    new_wards = []
    for r in new_rows:
        bnv_ma = str(r[1]).zfill(2)
        bnv_ten = r[2]
        ward_ma = str(r[7]).zfill(8)
        ward_ten = r[8]
        new_provinces[bnv_ma] = bnv_ten
        new_wards.append((ward_ma, ward_ten, bnv_ma))

    assert len(new_provinces) == 34
    assert len(new_wards) == 3321

    old_provinces = [(r[0], r[1]) for r in old_tinh_rows]
    assert len(old_provinces) == 34

    new_by_norm = defaultdict(list)
    for ma, ten in new_provinces.items():
        new_by_norm[norm_province(ten)].append(ma)

    old_to_new_ma = {}
    mapping_errors = []
    for old_ma, old_ten in old_provinces:
        if old_ma == "AG":
            old_to_new_ma[old_ma] = "32"  # quyet dinh nghiep vu #3 - An Giang -> BNV 32 (giong script truoc)
            continue
        key = norm_province(old_ten)
        candidates = new_by_norm.get(key, [])
        if len(candidates) == 1:
            old_to_new_ma[old_ma] = candidates[0]
        else:
            mapping_errors.append((old_ma, old_ten, key, candidates))

    if mapping_errors:
        raise SystemExit(f"DUNG: loi mapping tinh: {mapping_errors}")

    assert len(old_to_new_ma) == 34
    assert len(set(old_to_new_ma.values())) == 34

    # ---- Lookup old ward (5-digit cu) -> (ten, old_tinh_ma) tu backup phuong_xa cu ----
    old_ward_by_ma = {}
    for ma, ten, cap, parent_ma in old_pxa_rows:
        old_ward_by_ma[str(ma)] = (ten, str(parent_ma) if parent_ma is not None else None)

    # ---- new wards grouped by province (ma BNV moi), norm ten -> [new_ma] ----
    new_ward_by_province_norm = defaultdict(lambda: defaultdict(list))
    for ma, ten, parent_ma in new_wards:
        new_ward_by_province_norm[parent_ma][norm_ward(ten)].append(ma)

    matched = 0
    unmatched = []
    ambiguous = []
    cross_check_mismatch = []
    out_rows = []

    for row in so_gddt_rows:
        ma_don_vi, ten_don_vi, loai_don_vi, dia_ban_ma, don_vi_cha_ma = row

        # Cach 1: suy tinh cu tu ma_don_vi (dang "SO-GDDT-XX", giong format don_vi_cha_ma cua truong)
        old_tinh_ma_from_madonvi = str(ma_don_vi).replace("SO-GDDT-", "", 1)

        # Cach 2: tra nguoc dia_ban_ma (GSO cu) trong backup phuong_xa cu -> parent_ma (tinh cu)
        old_ward_ten, old_tinh_ma_from_diaban = old_ward_by_ma.get(str(dia_ban_ma), (None, None))

        if old_ward_ten is not None and old_tinh_ma_from_diaban != old_tinh_ma_from_madonvi:
            cross_check_mismatch.append(
                (ma_don_vi, ten_don_vi, dia_ban_ma, old_tinh_ma_from_madonvi, old_tinh_ma_from_diaban)
            )

        # Uu tien cach 2 (tra qua dia_danh cu) nhu nguoi dung yeu cau - khong phu thuoc parse ten
        old_tinh_ma = old_tinh_ma_from_diaban if old_tinh_ma_from_diaban is not None else old_tinh_ma_from_madonvi
        new_tinh_ma = old_to_new_ma.get(old_tinh_ma)

        new_dia_ban_ma = dia_ban_ma
        if old_ward_ten is None:
            unmatched.append((ma_don_vi, ten_don_vi, old_tinh_ma, dia_ban_ma, "khong tim thay dia_ban_ma trong du lieu phuong/xa cu (backup)"))
        elif new_tinh_ma is None:
            unmatched.append((ma_don_vi, ten_don_vi, old_tinh_ma, dia_ban_ma, "khong xac dinh duoc tinh moi"))
        else:
            key = norm_ward(old_ward_ten)
            candidates = new_ward_by_province_norm[new_tinh_ma].get(key, [])
            if len(candidates) == 1:
                new_dia_ban_ma = candidates[0]
                matched += 1
            elif len(candidates) == 0:
                unmatched.append((ma_don_vi, ten_don_vi, old_tinh_ma, dia_ban_ma, f"ten cu {old_ward_ten!r} (norm={key!r}) khong khop ten nao trong tinh moi {new_tinh_ma}"))
            else:
                ambiguous.append((ma_don_vi, ten_don_vi, old_tinh_ma, dia_ban_ma, old_ward_ten, key, candidates))

        out_rows.append((ma_don_vi, ten_don_vi, loai_don_vi, new_dia_ban_ma, don_vi_cha_ma))

    out_lines.append(f"Doi chieu 2 cach suy tinh cu (ma_don_vi vs dia_ban_ma->backup): mismatch = {len(cross_check_mismatch)}")
    for m in cross_check_mismatch:
        out_lines.append(f"  MISMATCH: {m}")

    out_lines.append(f"Tong dong: {len(so_gddt_rows)} | match tu dong: {matched} | khong match: {len(unmatched)} | mo ho: {len(ambiguous)}")

    if unmatched:
        out_lines.append("\n== KHONG MATCH (giu nguyen ma cu - CAN SOAT THU CONG) ==")
        for u in unmatched:
            out_lines.append(f"  {u}")

    if ambiguous:
        out_lines.append("\n== MO HO (nhieu ung vien - giu nguyen ma cu) ==")
        for a in ambiguous:
            out_lines.append(f"  {a}")

    if cross_check_mismatch:
        out_lines.append("\n[DUNG - KHONG GHI FILE] co mismatch giua 2 cach suy tinh cu, can soat lai logic truoc khi ghi.")
    else:
        assert len(out_rows) == len(so_gddt_rows) == 34, "Loi: so dong output khong khop 34"
        write_xlsx(SO_GDDT, ["ma_don_vi", "ten_don_vi", "loai_don_vi", "dia_ban_ma", "don_vi_cha_ma"], out_rows)
        out_lines.append(f"\n[write] {SO_GDDT} -> {len(out_rows)} dong (da ghi de; {matched} dong remap, {len(unmatched)+len(ambiguous)} dong giu nguyen ma cu vi khong tu khop duoc - xem danh sach tren)")

    report_path = ROOT / "scripts" / "_report_remap_so_gddt.md"
    report_path.write_text("\n".join(out_lines), encoding="utf-8")

    for line in out_lines:
        print(line)


if __name__ == "__main__":
    main()
